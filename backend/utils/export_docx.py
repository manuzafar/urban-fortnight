"""
DOCX export utility for generating inception pack Word documents.

Uses python-docx for document generation.
"""

from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE


def generate_docx(pack: dict[str, Any], section: str | None = None) -> bytes:
    """
    Generate a DOCX file from an inception pack.

    Args:
        pack: The inception pack dictionary containing all sections.
        section: Optional section name to export only that section.
                 If None, exports the full pack.

    Returns:
        bytes: The generated DOCX content.
    """
    doc = Document()

    # Set up styles
    _setup_styles(doc)

    # Get product name
    product_name = "Product"
    if pack.get("executive_summary") and pack["executive_summary"].get("product_name"):
        product_name = pack["executive_summary"]["product_name"]

    # Title page (only for full pack)
    if not section:
        _add_title_page(doc, pack, product_name)

    # Define sections to render
    all_sections = [
        ("executive_summary", "1. Executive Summary"),
        ("customer_research", "2. Customer Research"),
        ("business_case", "3. Business Case"),
        ("product_requirements_document", "4. Product Requirements Document"),
        ("technical_architecture", "5. Technical Architecture"),
        ("legal_regulatory_review", "6. Legal & Regulatory Review"),
        ("quality_assessment", "7. Quality Assessment"),
    ]

    sections_to_render = [(s, t) for s, t in all_sections if not section or s == section]

    for sec_key, sec_title in sections_to_render:
        if sec_key not in pack or pack[sec_key] is None:
            continue

        doc.add_heading(sec_title, level=1)
        _render_section(doc, pack[sec_key], sec_key)

        if not section:
            doc.add_page_break()

    # Save to bytes
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _setup_styles(doc: Document) -> None:
    """Set up document styles."""
    # Modify default styles
    style = doc.styles["Heading 1"]
    style.font.size = Pt(18)
    style.font.color.rgb = RGBColor(0x4a, 0x4a, 0x6a)

    style = doc.styles["Heading 2"]
    style.font.size = Pt(14)
    style.font.color.rgb = RGBColor(0x66, 0x7e, 0xea)

    style = doc.styles["Heading 3"]
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(0x4a, 0x4a, 0x6a)


def _add_title_page(doc: Document, pack: dict[str, Any], product_name: str) -> None:
    """Add a title page to the document."""
    # Add spacing
    for _ in range(6):
        doc.add_paragraph()

    # Title
    title = doc.add_heading(product_name, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Tagline
    if pack.get("executive_summary") and pack["executive_summary"].get("tagline"):
        tagline = doc.add_paragraph(pack["executive_summary"]["tagline"])
        tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Subtitle
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Inception Pack")
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x4a, 0x4a, 0x6a)

    # Metadata
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()

    if pack.get("metadata"):
        meta = doc.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta.add_run(f"Generated: {pack['metadata'].get('generated_at', 'Unknown')}\n")
        meta.add_run(f"Session: {pack['metadata'].get('session_id', 'Unknown')[:8]}...")

    doc.add_page_break()


def _render_section(doc: Document, data: dict[str, Any], section_name: str) -> None:
    """Render a section's content."""
    if section_name == "executive_summary":
        _render_executive_summary(doc, data)
    elif section_name == "customer_research":
        _render_customer_research(doc, data)
    elif section_name == "business_case":
        _render_business_case(doc, data)
    elif section_name == "product_requirements_document":
        _render_prd(doc, data)
    elif section_name == "technical_architecture":
        _render_technical_architecture(doc, data)
    elif section_name == "legal_regulatory_review":
        _render_legal_regulatory(doc, data)
    elif section_name == "quality_assessment":
        _render_quality_assessment(doc, data)
    else:
        _render_generic(doc, data)


def _render_executive_summary(doc: Document, data: dict[str, Any]) -> None:
    """Render executive summary section."""
    if data.get("product_name"):
        doc.add_heading(data["product_name"], level=2)
        if data.get("tagline"):
            p = doc.add_paragraph()
            run = p.add_run(data["tagline"])
            run.italic = True

    if data.get("problem_statement"):
        doc.add_heading("Problem Statement", level=3)
        doc.add_paragraph(data["problem_statement"])

    if data.get("solution_overview"):
        doc.add_heading("Solution Overview", level=3)
        doc.add_paragraph(data["solution_overview"])

    if data.get("target_audience"):
        doc.add_heading("Target Audience", level=3)
        doc.add_paragraph(data["target_audience"])

    if data.get("unique_value_proposition"):
        doc.add_heading("Unique Value Proposition", level=3)
        doc.add_paragraph(data["unique_value_proposition"])

    if data.get("key_metrics"):
        doc.add_heading("Key Metrics", level=3)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "Metric"
        hdr[1].text = "Target"
        hdr[2].text = "Rationale"

        for metric in data["key_metrics"]:
            row = table.add_row().cells
            row[0].text = str(metric.get("name") or metric.get("metric") or metric)
            row[1].text = str(metric.get("target") or metric.get("value") or "TBD")
            row[2].text = str(metric.get("rationale") or metric.get("description") or "-")

    if data.get("key_risks"):
        doc.add_heading("Key Risks", level=3)
        for risk in data["key_risks"][:5]:
            risk_text = risk.get("risk") if isinstance(risk, dict) else str(risk)
            doc.add_paragraph(risk_text, style="List Bullet")

    if data.get("recommendation"):
        doc.add_heading("Recommendation", level=3)
        doc.add_paragraph(data["recommendation"])


def _render_customer_research(doc: Document, data: dict[str, Any]) -> None:
    """Render customer research section."""
    if data.get("target_customer"):
        doc.add_heading("Target Customer", level=2)
        doc.add_paragraph(data["target_customer"])

    if data.get("pain_signals"):
        doc.add_heading("Pain Signals", level=2)
        for signal in data["pain_signals"][:8]:
            if isinstance(signal, dict):
                tier = signal.get("evidence_tier", "E3")
                text = signal.get("signal") or signal.get("description") or str(signal)
                doc.add_paragraph(f"[{tier}] {text}", style="List Bullet")
            else:
                doc.add_paragraph(str(signal), style="List Bullet")

    if data.get("market_hypotheses"):
        doc.add_heading("Market Hypotheses", level=2)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "Hypothesis"
        hdr[1].text = "Validation"
        hdr[2].text = "Evidence Tier"

        for hyp in data["market_hypotheses"][:6]:
            row = table.add_row().cells
            if isinstance(hyp, dict):
                row[0].text = str(hyp.get("hypothesis") or hyp)
                row[1].text = str(hyp.get("validation") or "-")
                row[2].text = str(hyp.get("evidence_tier") or "E3")
            else:
                row[0].text = str(hyp)
                row[1].text = "-"
                row[2].text = "E3"

    if data.get("uncomfortable_insights"):
        doc.add_heading("Uncomfortable Insights", level=2)
        for insight in data["uncomfortable_insights"][:4]:
            doc.add_paragraph(str(insight), style="List Bullet")

    if data.get("competitive_landscape"):
        doc.add_heading("Competitive Landscape", level=2)
        doc.add_paragraph(data["competitive_landscape"])


def _render_business_case(doc: Document, data: dict[str, Any]) -> None:
    """Render business case section."""
    if data.get("lean_canvas"):
        doc.add_heading("Lean Canvas", level=2)
        canvas = data["lean_canvas"]
        table = doc.add_table(rows=8, cols=2)
        table.style = "Table Grid"

        fields = [
            ("Problem", canvas.get("problem")),
            ("Solution", canvas.get("solution")),
            ("Unique Value Proposition", canvas.get("unique_value_proposition")),
            ("Unfair Advantage", canvas.get("unfair_advantage")),
            ("Customer Segments", canvas.get("customer_segments")),
            ("Channels", canvas.get("channels")),
            ("Revenue Streams", canvas.get("revenue_streams")),
            ("Cost Structure", canvas.get("cost_structure")),
        ]

        for i, (label, value) in enumerate(fields):
            table.rows[i].cells[0].text = label
            table.rows[i].cells[1].text = str(value or "-")

    if data.get("market_sizing"):
        doc.add_heading("Market Sizing", level=2)
        sizing = data["market_sizing"]
        table = doc.add_table(rows=3, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "TAM (Total Addressable Market)"
        table.rows[0].cells[1].text = str(sizing.get("tam") or "-")
        table.rows[1].cells[0].text = "SAM (Serviceable Addressable Market)"
        table.rows[1].cells[1].text = str(sizing.get("sam") or "-")
        table.rows[2].cells[0].text = "SOM (Serviceable Obtainable Market)"
        table.rows[2].cells[1].text = str(sizing.get("som") or "-")

    if data.get("financial_projections"):
        doc.add_heading("Financial Projections", level=2)
        fp = data["financial_projections"]
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "Period"
        hdr[1].text = "Revenue"
        hdr[2].text = "Costs"
        hdr[3].text = "Profit"

        for year, label in [("year_1", "Year 1"), ("year_2", "Year 2"), ("year_3", "Year 3")]:
            if year in fp and fp[year]:
                row = table.add_row().cells
                row[0].text = label
                row[1].text = str(fp[year].get("revenue") or "TBD")
                row[2].text = str(fp[year].get("costs") or "TBD")
                row[3].text = str(fp[year].get("profit") or "TBD")

    if data.get("gtm_strategy"):
        doc.add_heading("Go-to-Market Strategy", level=2)
        doc.add_paragraph(data["gtm_strategy"])

    if data.get("pricing_strategy"):
        doc.add_heading("Pricing Strategy", level=2)
        doc.add_paragraph(data["pricing_strategy"])


def _render_prd(doc: Document, data: dict[str, Any]) -> None:
    """Render PRD section."""
    if data.get("product_vision"):
        doc.add_heading("Product Vision", level=2)
        doc.add_paragraph(data["product_vision"])

    if data.get("goals"):
        doc.add_heading("Goals", level=2)
        for goal in data["goals"]:
            doc.add_paragraph(str(goal), style="List Bullet")

    if data.get("epics"):
        doc.add_heading("Epics", level=2)
        for i, epic in enumerate(data["epics"][:5], 1):
            epic_id = epic.get("id") or f"E{i}"
            epic_title = epic.get("title") or epic.get("name") or "Epic"
            doc.add_heading(f"{epic_id}: {epic_title}", level=3)

            if epic.get("description"):
                doc.add_paragraph(epic["description"])

            if epic.get("priority"):
                doc.add_paragraph(f"Priority: {epic['priority']}")

            if epic.get("user_stories"):
                doc.add_paragraph(f"User Stories: {len(epic['user_stories'])} stories")

    if data.get("functional_requirements"):
        doc.add_heading("Functional Requirements", level=2)
        for req in data["functional_requirements"][:10]:
            req_text = req.get("requirement") if isinstance(req, dict) else str(req)
            if isinstance(req, dict):
                req_text = req.get("requirement") or req.get("description") or str(req)
            doc.add_paragraph(req_text, style="List Bullet")

    if data.get("non_functional_requirements"):
        doc.add_heading("Non-Functional Requirements", level=2)
        for req in data["non_functional_requirements"][:8]:
            req_text = req.get("requirement") if isinstance(req, dict) else str(req)
            if isinstance(req, dict):
                req_text = req.get("requirement") or req.get("description") or str(req)
            doc.add_paragraph(req_text, style="List Bullet")

    if data.get("mvp_scope"):
        doc.add_heading("MVP Scope", level=2)
        doc.add_paragraph(data["mvp_scope"])

    if data.get("success_criteria"):
        doc.add_heading("Success Criteria", level=2)
        for criteria in data["success_criteria"]:
            doc.add_paragraph(str(criteria), style="List Bullet")


def _render_technical_architecture(doc: Document, data: dict[str, Any]) -> None:
    """Render technical architecture section."""
    if data.get("architecture_overview"):
        doc.add_heading("Architecture Overview", level=2)
        doc.add_paragraph(data["architecture_overview"])

    if data.get("technology_stack"):
        doc.add_heading("Technology Stack", level=2)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "Category"
        hdr[1].text = "Technologies"

        tech_stack = data["technology_stack"]
        if isinstance(tech_stack, dict):
            for category, techs in tech_stack.items():
                row = table.add_row().cells
                row[0].text = category.replace("_", " ").title()
                if isinstance(techs, list):
                    row[1].text = ", ".join(str(t) for t in techs)
                else:
                    row[1].text = str(techs)
        elif isinstance(tech_stack, list):
            for item in tech_stack:
                row = table.add_row().cells
                if isinstance(item, dict):
                    row[0].text = str(item.get("category") or item.get("name") or "Technology")
                    techs = item.get("technologies") or item.get("tech") or item.get("value") or ""
                    if isinstance(techs, list):
                        row[1].text = ", ".join(str(t) for t in techs)
                    else:
                        row[1].text = str(techs)
                else:
                    row[0].text = "Technology"
                    row[1].text = str(item)

    if data.get("key_components"):
        doc.add_heading("Key Components", level=2)
        for comp in data["key_components"][:6]:
            doc.add_heading(comp.get("name", "Component"), level=3)
            if comp.get("description"):
                doc.add_paragraph(comp["description"])
            if comp.get("technologies"):
                techs = comp["technologies"]
                if isinstance(techs, list):
                    techs = ", ".join(str(t) for t in techs)
                doc.add_paragraph(f"Technologies: {techs}")

    if data.get("integration_points"):
        doc.add_heading("Integration Points", level=2)
        for point in data["integration_points"]:
            if isinstance(point, dict):
                name = point.get("name", "Integration")
                desc = point.get("description", "")
                doc.add_paragraph(f"{name}: {desc}", style="List Bullet")
            else:
                doc.add_paragraph(str(point), style="List Bullet")

    if data.get("security_considerations"):
        doc.add_heading("Security Considerations", level=2)
        for item in data["security_considerations"]:
            doc.add_paragraph(str(item), style="List Bullet")

    if data.get("scalability_approach"):
        doc.add_heading("Scalability Approach", level=2)
        doc.add_paragraph(data["scalability_approach"])


def _render_legal_regulatory(doc: Document, data: dict[str, Any]) -> None:
    """Render legal & regulatory section."""
    if data.get("applicable_regulations"):
        doc.add_heading("Applicable Regulations", level=2)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "Regulation"
        hdr[1].text = "Requirements"
        hdr[2].text = "Impact"

        for reg in data["applicable_regulations"]:
            row = table.add_row().cells
            if isinstance(reg, dict):
                row[0].text = str(reg.get("name") or reg.get("regulation") or reg)
                row[1].text = str(reg.get("requirements") or "-")
                row[2].text = str(reg.get("impact") or "-")
            else:
                row[0].text = str(reg)
                row[1].text = "-"
                row[2].text = "-"

    if data.get("data_protection"):
        doc.add_heading("Data Protection", level=2)
        dp = data["data_protection"]
        if isinstance(dp, dict):
            text = dp.get("summary") or dp.get("requirements") or str(dp)
        else:
            text = str(dp)
        doc.add_paragraph(text)

    if data.get("risks"):
        doc.add_heading("Legal Risks", level=2)
        for risk in data["risks"][:5]:
            if isinstance(risk, dict):
                risk_name = risk.get("risk") or risk.get("title") or "Risk"
                mitigation = risk.get("mitigation") or risk.get("description") or ""
                doc.add_paragraph(f"{risk_name}: {mitigation}", style="List Bullet")
            else:
                doc.add_paragraph(str(risk), style="List Bullet")

    if data.get("ip_considerations"):
        doc.add_heading("Intellectual Property Considerations", level=2)
        doc.add_paragraph(data["ip_considerations"])

    if data.get("compliance_roadmap"):
        doc.add_heading("Compliance Roadmap", level=2)
        doc.add_paragraph(data["compliance_roadmap"])


def _render_quality_assessment(doc: Document, data: dict[str, Any]) -> None:
    """Render quality assessment section."""
    if data.get("overall_score") is not None:
        doc.add_heading("Overall Quality Score", level=2)
        score_pct = round(data["overall_score"] * 100)
        p = doc.add_paragraph()
        run = p.add_run(f"{score_pct}%")
        run.bold = True
        run.font.size = Pt(24)
        run.font.color.rgb = RGBColor(0x66, 0x7e, 0xea)

    if data.get("section_scores"):
        doc.add_heading("Section Scores", level=2)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "Section"
        hdr[1].text = "Score"

        section_scores = data["section_scores"]
        if isinstance(section_scores, dict):
            for section_name, score in section_scores.items():
                row = table.add_row().cells
                row[0].text = section_name.replace("_", " ").title()
                row[1].text = f"{round(score * 100)}%"
        elif isinstance(section_scores, list):
            for item in section_scores:
                row = table.add_row().cells
                if isinstance(item, dict):
                    row[0].text = str(item.get("section") or item.get("name") or "Section").replace("_", " ").title()
                    score = item.get("score") or item.get("value") or 0
                    row[1].text = f"{round(float(score) * 100)}%"
                else:
                    row[0].text = "Section"
                    row[1].text = str(item)

    if data.get("strengths"):
        doc.add_heading("Strengths", level=2)
        for strength in data["strengths"]:
            doc.add_paragraph(str(strength), style="List Bullet")

    if data.get("critical_gaps"):
        doc.add_heading("Areas Needing Attention", level=2)
        for gap in data["critical_gaps"]:
            doc.add_paragraph(str(gap), style="List Bullet")

    if data.get("recommendations"):
        doc.add_heading("Recommendations", level=2)
        for rec in data["recommendations"]:
            doc.add_paragraph(str(rec), style="List Bullet")


def _render_generic(doc: Document, data: dict[str, Any]) -> None:
    """Render any section generically."""
    for key, value in data.items():
        if value is None:
            continue

        heading = key.replace("_", " ").title()
        doc.add_heading(heading, level=2)

        if isinstance(value, str):
            doc.add_paragraph(value)
        elif isinstance(value, list):
            for item in value:
                if isinstance(item, dict):
                    for k, v in item.items():
                        doc.add_paragraph(f"{k}: {v}", style="List Bullet")
                else:
                    doc.add_paragraph(str(item), style="List Bullet")
        elif isinstance(value, dict):
            table = doc.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            for k, v in value.items():
                row = table.add_row().cells
                row[0].text = str(k)
                row[1].text = str(v) if v else ""


def get_docx_filename(session_id: str, section: str | None = None) -> str:
    """
    Generate a filename for the DOCX export.

    Args:
        session_id: The session ID.
        section: Optional section name.

    Returns:
        str: The suggested filename.
    """
    short_id = session_id[:8]
    if section:
        return f"{section.replace('_', '-')}-{short_id}.docx"
    return f"inception-pack-{short_id}.docx"
